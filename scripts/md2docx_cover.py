import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

src, dst = sys.argv[1], sys.argv[2]
lines = open(src, encoding="utf-8").read().split("\n")
d = Document()
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(__import__("docx").oxml.ns.qn("w:eastAsia"), "宋体")
for s in d.sections:
    s.top_margin = s.bottom_margin = Cm(2.54); s.left_margin = s.right_margin = Cm(2.54)

def add(txt, bold=False, italic=False, align=None):
    p = d.add_paragraph(); r = p.add_run(txt); r.bold = bold; r.italic = italic
    r.font.color.rgb = RGBColor(0, 0, 0)
    p.paragraph_format.space_after = Pt(6)
    if align: p.alignment = align
    return p

for ln in lines:
    t = ln.strip()
    if not t: continue
    if t.startswith("# "): continue
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"\*(.+?)\*", r"\1", t)
    add(t, bold=False)
d.save(dst)
print("Cover letter DOCX:", dst)
